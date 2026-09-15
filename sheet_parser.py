import io

import openpyxl

INFO_LABELS = {
    "이름": "name",
    "직업": "occupation",
    "나이": "age",
    "성별": "sex",
    "거주지": "residence",
    "출생지": "birthplace",
}

ATTRIBUTE_LABELS = {
    "근력": "str",
    "민첩": "dex",
    "정신": "pow",
    "건강": "con",
    "외모": "app",
    "교육": "edu",
    "크기": "siz",
    "지능": "int",
}

_INT_FIELDS = set(ATTRIBUTE_LABELS.values()) | {"age"}

SKILL_NAMES = {
    "감정", "고고학", "관찰력", "근접전투(격투)", "기계수리", "도약", "듣기", "말주변",
    "매혹", "법률", "변장", "사격(권총)", "사격(라이플/샷건)", "설득", "손놀림", "수영",
    "승마", "심리학", "언어(모국어)", "역사", "열쇠공", "오르기", "오컬트", "위협",
    "은밀행동", "응급치료", "의료", "인류학", "자동차 운전", "자료조사", "자연", "재력",
    "전기수리", "정신분석", "중장비 조작", "추적", "크툴루 신화", "투척", "항법", "회계",
    "회피",
}


def _build_label_index(ws) -> dict:
    index = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                label = cell.value.strip()
                if label and label not in index:
                    index[label] = cell
    return index


def _adjacent_value(ws, cell):
    return ws.cell(row=cell.row, column=cell.column + 1).value


def _to_int(value, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"'{label}' 값이 숫자가 아닙니다: {value!r}")
    return int(value)


def _compute_mov(str_: int, dex: int, siz: int, age: int) -> int:
    if str_ < siz and dex < siz:
        mov = 7
    elif str_ > siz and dex > siz:
        mov = 9
    else:
        mov = 8
    age_penalty = max(0, (min(age, 89) - 30) // 10)
    return mov - age_penalty


def _compute_hp_max(con: int, siz: int) -> int:
    return (con + siz) // 10


def _compute_mp_max(pow_: int) -> int:
    return pow_ // 5


def _parse_skills(ws, labels: dict) -> dict:
    skills = {}
    for name in SKILL_NAMES:
        cell = labels.get(name)
        if cell is None:
            continue
        value = _adjacent_value(ws, cell)
        if isinstance(value, (int, float)) and not isinstance(value, bool) and value:
            skills[name] = int(value)
    return skills


def parse_character_sheet(file_bytes: bytes) -> dict:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ValueError("xlsx 파일이 아니거나 파일이 손상되었습니다.") from exc
    ws = wb.active
    labels = _build_label_index(ws)

    result = {}
    for label, field in {**INFO_LABELS, **ATTRIBUTE_LABELS}.items():
        cell = labels.get(label)
        if cell is None:
            raise ValueError(f"'{label}' 항목을 시트에서 찾을 수 없습니다.")
        value = _adjacent_value(ws, cell)
        result[field] = _to_int(value, label) if field in _INT_FIELDS else value

    if not result.get("name") or not str(result["name"]).strip():
        raise ValueError("'이름' 값이 비어 있습니다.")

    result["mov"] = _compute_mov(result["str"], result["dex"], result["siz"], result["age"])
    result["hp_max"] = result["hp_current"] = _compute_hp_max(result["con"], result["siz"])
    result["mp_max"] = result["mp_current"] = _compute_mp_max(result["pow"])
    result["san_starting"] = result["san_current"] = result["pow"]
    result["skills"] = _parse_skills(ws, labels)
    return result
