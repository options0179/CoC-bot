import io

import openpyxl
import pytest

from sheet_parser import _compute_hp_max, _compute_mov, _compute_mp_max, parse_character_sheet

_BASE_FIELDS = {
    "이름": "탐사자",
    "직업": "사립탐정",
    "나이": 32,
    "성별": "여",
    "거주지": "서울",
    "출생지": "부산",
    "근력": 50,
    "민첩": 60,
    "정신": 55,
    "건강": 65,
    "외모": 45,
    "교육": 70,
    "크기": 50,
    "지능": 80,
}


def _build_workbook(fields: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row, (label, value) in enumerate(fields.items(), start=1):
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_character_sheet_extracts_info_fields():
    result = parse_character_sheet(_build_workbook(_BASE_FIELDS))
    assert result["name"] == "탐사자"
    assert result["occupation"] == "사립탐정"
    assert result["age"] == 32
    assert result["sex"] == "여"
    assert result["residence"] == "서울"
    assert result["birthplace"] == "부산"


def test_parse_character_sheet_extracts_attributes():
    result = parse_character_sheet(_build_workbook(_BASE_FIELDS))
    assert result["str"] == 50
    assert result["dex"] == 60
    assert result["pow"] == 55
    assert result["con"] == 65
    assert result["app"] == 45
    assert result["edu"] == 70
    assert result["siz"] == 50
    assert result["int"] == 80
    assert result["mov"] == 8


def test_parse_character_sheet_computes_hp_mp_san():
    result = parse_character_sheet(_build_workbook(_BASE_FIELDS))
    assert result["hp_max"] == result["hp_current"] == 11  # (65+50)//10
    assert result["mp_max"] == result["mp_current"] == 11  # 55//5
    assert result["san_starting"] == result["san_current"] == 55  # == POW


def test_parse_character_sheet_missing_label_raises():
    fields = dict(_BASE_FIELDS)
    del fields["직업"]
    with pytest.raises(ValueError, match="직업"):
        parse_character_sheet(_build_workbook(fields))


def test_parse_character_sheet_non_numeric_attribute_raises():
    fields = dict(_BASE_FIELDS, 근력="abc")
    with pytest.raises(ValueError, match="근력"):
        parse_character_sheet(_build_workbook(fields))


def test_parse_character_sheet_extracts_skills():
    fields = dict(_BASE_FIELDS, 회계=5, 심리학=10, 회피=30)
    result = parse_character_sheet(_build_workbook(fields))
    assert result["skills"] == {"회계": 5, "심리학": 10, "회피": 30}


def test_parse_character_sheet_ignores_zero_valued_skills():
    fields = dict(_BASE_FIELDS, 회계=0, 심리학=10)
    result = parse_character_sheet(_build_workbook(fields))
    assert result["skills"] == {"심리학": 10}


def test_parse_character_sheet_skills_default_empty_when_none_present():
    result = parse_character_sheet(_build_workbook(_BASE_FIELDS))
    assert result["skills"] == {}


def test_parse_character_sheet_non_xlsx_raises_value_error():
    with pytest.raises(ValueError, match="xlsx"):
        parse_character_sheet(b"not an xlsx file at all")


def test_parse_character_sheet_empty_name_raises():
    fields = dict(_BASE_FIELDS, 이름="   ")
    with pytest.raises(ValueError, match="이름"):
        parse_character_sheet(_build_workbook(fields))


def test_compute_mov_both_below_size_is_slow():
    assert _compute_mov(str_=30, dex=30, siz=50, age=25) == 7


def test_compute_mov_both_above_size_is_fast():
    assert _compute_mov(str_=60, dex=60, siz=50, age=25) == 9


def test_compute_mov_mixed_is_average():
    assert _compute_mov(str_=30, dex=60, siz=50, age=25) == 8


def test_compute_mov_applies_age_penalty():
    assert _compute_mov(str_=60, dex=60, siz=50, age=45) == 8
    assert _compute_mov(str_=60, dex=60, siz=50, age=85) == 4


def test_compute_hp_max_rounds_down():
    assert _compute_hp_max(con=65, siz=50) == 11
    assert _compute_hp_max(con=64, siz=50) == 11
    assert _compute_hp_max(con=64, siz=45) == 10


def test_compute_mp_max_rounds_down():
    assert _compute_mp_max(pow_=55) == 11
    assert _compute_mp_max(pow_=59) == 11
    assert _compute_mp_max(pow_=60) == 12
